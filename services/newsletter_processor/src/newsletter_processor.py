import json
import logging
import os
from datetime import datetime, timedelta
from io import BytesIO

import boto3
import docx
from docx import Document
from botocore.exceptions import ClientError

from constants import *

logger = logging.getLogger(__name__)
logging.getLogger().setLevel(logging.INFO)

ssm = boto3.client('ssm')

def get_config():
    """Get configuration from SSM Parameter Store"""
    config_path = os.getenv("CONFIG_PATH")
    try:
        # First get list of sets
        sets_param = ssm.get_parameter(Name=f"{config_path}/arxiv/sets")
        set_names = json.loads(sets_param['Parameter']['Value'])
        
        # Get global params first
        global_params = ssm.get_parameters(
            Names=[
                f"{config_path}/arxiv/s3_bucket",
                f"{config_path}/arxiv/dynamodb_table"
            ]
        )
        
        # Initialize config structure
        config = {}
        for param in global_params['Parameters']:
            name = param['Name'].split('/')[-1]
            config[name] = param['Value']
        
        # Initialize sets config
        config["sets"] = {}
        
        # Get params for each set
        for set_name in set_names:
            set_params = ssm.get_parameters(
                Names=[
                    f"{config_path}/arxiv/{set_name}/categories",
                    f"{config_path}/arxiv/{set_name}/back_date"
                ]
            )
            
            set_config = {}
            for param in set_params['Parameters']:
                name = param['Name'].split('/')[-1]
                if name == 'categories':
                    set_config[name] = param['Value'].split(',')
                elif name == 'back_date':
                    set_config[name] = int(param['Value'])
            
            config["sets"][set_name] = set_config
            
        # For backward compatibility, default to the first set if available
        if set_names and set_names[0] in config["sets"]:
            default_set = set_names[0]
            config["categories"] = config["sets"][default_set]["categories"]
            config["back_date"] = config["sets"][default_set]["back_date"]
            config["set"] = default_set
            
        return config
    except ClientError as e:
        logger.error(f"Error fetching config: {e}")
        raise

def upload_to_s3(file_data: BytesIO, key: str, s3_bucket: str, s3_client):
    """Upload file to S3"""
    try:
        s3_client.upload_fileobj(file_data, s3_bucket, key)
        logger.info(f"Uploaded {key} to S3")
    except ClientError as e:
        logger.error(f"Error uploading to S3: {e}")
        raise

def add_hyperlink(paragraph, text, url):
    """Add a hyperlink to a paragraph"""
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    
    hyperlink = docx.oxml.shared.OxmlElement('w:hyperlink')
    hyperlink.set(docx.oxml.shared.qn('r:id'), r_id)
    
    new_run = docx.oxml.shared.OxmlElement('w:r')
    rPr = docx.oxml.shared.OxmlElement('w:rPr')
    
    rStyle = docx.oxml.shared.OxmlElement('w:rStyle')
    rStyle.set(docx.oxml.shared.qn('w:val'), 'Hyperlink')
    rPr.append(rStyle)
    
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    
    paragraph._p.append(hyperlink)
    
    return hyperlink

def create_research_summary(records: list, date: str, categories: list, s3_bucket: str, s3_client, set_name: str) -> dict:
    """Creates research summary documents and returns file info"""
    summary_files = {}
    
    for category in categories:
        doc = Document()
        doc.add_heading(f"arXiv {set_name}/{category} Research Summaries - {date}", 0)
        
        # Filter for records with this category
        category_papers = []
        for record in records:
            # Date comparison logic
            start_date = datetime.strptime(date, "%Y-%m-%d")
            end_date = start_date + timedelta(days=1)
            record_date = datetime.strptime(record["date"], "%Y-%m-%d")

            if record["primary_category"] == category and start_date <= record_date < end_date:
                # Store paper info for metadata
                category_papers.append(record)
                
                # Add title with link
                title_para = doc.add_paragraph()
                add_hyperlink(title_para, record["title"], record["abstract_url"])
                
                # Add PDF link
                pdf_link = record["abstract_url"].replace("abs", "pdf")
                title_para.add_run(" [")
                add_hyperlink(title_para, "PDF", pdf_link)
                title_para.add_run("]")
                
                # Add authors
                authors = [f"{author['first_name']} {author['last_name']}" for author in record["authors"]]
                doc.add_paragraph(f"Authors: {', '.join(authors)}")
                
                # Add abstract
                abstract = latex_to_human_readable(record["abstract"])
                doc.add_paragraph(abstract)
                
                # Add spacing
                doc.add_paragraph()

        if category_papers:
            # Save to memory
            docx_buffer = BytesIO()
            doc.save(docx_buffer)
            docx_buffer.seek(0)
            
            # Updated S3 path to include set
            s3_key = f"newsletters/{date}/{set_name}/{category}_research_summary.docx"
            upload_to_s3(docx_buffer, s3_key, s3_bucket, s3_client)
            
            summary_files[category] = {
                "s3_key": s3_key,
                "papers": category_papers
            }
            
            logging.info(f"Created and uploaded summary for {category}")

    return summary_files

def main():    
    logger.info("Starting newsletter processor")
    exit()
    # Create and upload summary documents
    summary_files = create_research_summary(
        all_records, 
        date, 
        categories, 
        config["s3_bucket"], 
        s3_client,
        set_name
    )
    
    if summary_files:
        logging.info(f"Successfully processed {len(all_records)} papers for {set_name}/{date}")
    else:
        logging.warning(f"No papers in selected categories for {set_name}/{date}")